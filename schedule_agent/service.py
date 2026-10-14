from __future__ import annotations

import json
import shutil
import traceback
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .analyzer import analyze_bundle
from .excel_io import WorkbookInputError, create_template, export_result, read_bundle
from .pdf_report import export_pdf
from .store import Store
from .wbs_rules import generate_candidates


class ScheduleService:
    def __init__(self, root: Path):
        self.root = root
        self.data_dir = root / "data"
        self.generated_dir = root / "generated"
        self.assets_dir = root / "assets"
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.generated_dir.mkdir(parents=True, exist_ok=True)
        self.store = Store(self.data_dir / "schedule_agent.sqlite3")
        self.template_path = self.generated_dir / "市政道路进度计划标准模板.xlsx"
        self.sample_path = self.generated_dir / "市政道路进度计划示例.xlsx"
        if not self.template_path.exists():
            create_template(self.template_path, sample=False)
        if not self.sample_path.exists():
            create_template(self.sample_path, sample=True)

    def analyze(self, content: bytes, filename: str, actor: str = "计划工程师") -> dict:
        try:
            bundle = read_bundle(content, filename)
            self._normalize_bundle(bundle)
        except Exception as exc:
            self._record_import_error(filename, "读取与字段识别", exc)
            if isinstance(exc, ValueError):
                raise
            raise WorkbookInputError("读取进度表时发生异常，请检查合并表头、日期和编码字段后重试。") from exc
        candidate_notes = generate_candidates(bundle)
        run_id = "RUN-" + uuid.uuid4().hex[:10].upper()
        version = self.store.next_version(bundle.project.project_id)
        try:
            result = analyze_bundle(bundle, run_id, version)
        except Exception as exc:
            self._record_import_error(filename, "排程与偏差计算", exc)
            if isinstance(exc, TypeError):
                raise WorkbookInputError(
                    "进度表中存在无法比较的混合数据类型。请检查 WBS/任务编码、工期和日期列；编码列建议统一设为文本。"
                ) from exc
            raise
        self._apply_candidate_issues(result, candidate_notes)
        self._apply_import_issues(result, bundle)
        run_dir = self.generated_dir / run_id
        run_dir.mkdir(parents=True, exist_ok=False)
        input_path = run_dir / "input.xlsx"
        input_path.write_bytes(content)
        self.store.add_run(run_id, bundle.project.project_id, version, result.formal_eligible, input_path, filename, bundle.input_hash, actor)
        try:
            self._write_outputs(result, run_dir)
        except Exception as exc:
            self._record_import_error(filename, "Excel/PDF 交付物生成", exc)
            raise WorkbookInputError("分析已完成，但生成 Excel/PDF 时发现不兼容数据，请检查异常长文本、日期和公式单元格。") from exc
        return self._payload(result)

    def analyze_sample(self) -> dict:
        return self.analyze(self.sample_path.read_bytes(), self.sample_path.name, "示例用户")

    def get_result(self, run_id: str) -> dict:
        metadata = self.store.get_run(run_id)
        if not metadata:
            raise ValueError("找不到该运行")
        content = Path(metadata["input_path"]).read_bytes()
        bundle = read_bundle(content, metadata["input_filename"])
        self._normalize_bundle(bundle)
        candidate_notes = generate_candidates(bundle)
        import_confirmed = bool(metadata.get("import_confirmed"))
        if import_confirmed:
            self._confirm_auto_content(bundle)
        adjustments = self.store.get_manual_adjustments(run_id)
        self._apply_manual_adjustments(bundle, adjustments)
        result = analyze_bundle(bundle, run_id, metadata["version"])
        self._apply_candidate_issues(result, candidate_notes, import_confirmed)
        self._apply_import_issues(result, bundle, import_confirmed)
        result.formal_eligible = bool(metadata["formal_eligible"])
        result.state = metadata["state"]
        self._append_manual_traceability(result, adjustments)
        payload = self._payload(result, import_confirmed, adjustments)
        payload["audit"] = self.store.audit(run_id)
        return payload

    def confirm_import(self, run_id: str, actor: str, acknowledged: bool) -> dict:
        if not acknowledged:
            raise ValueError("请先确认自动识别的字段映射、默认值和生成任务")
        metadata = self.store.get_run(run_id)
        if not metadata:
            raise ValueError("找不到该运行")
        preview = self.get_result(run_id)
        if not preview["requires_import_confirmation"]:
            raise ValueError("当前运行没有待确认的自动导入内容")
        comment = "已确认自动识别的字段映射、默认值、生成 WBS 及施工逻辑，进入人工复核流程。"
        self.store.confirm_import(run_id, actor, comment)
        result = self._recalculate(run_id, use_stored_eligibility=False)
        self.store.update_eligibility(run_id, result.formal_eligible)
        self._write_outputs(result, self.generated_dir / run_id)
        return self.get_result(run_id)

    def archive_plan(self, run_id: str, actor: str) -> dict:
        return self.store.archive_plan(run_id, actor)

    def adjust_progress(
        self,
        run_id: str,
        wbs: str,
        actual_percent: float,
        remaining_duration: int,
        reason: str,
        action_type: str,
        note: str,
        actor: str,
    ) -> dict:
        metadata = self.store.get_run(run_id)
        if not metadata:
            raise ValueError("找不到该运行")
        content = Path(metadata["input_path"]).read_bytes()
        bundle = read_bundle(content, metadata["input_filename"])
        self._normalize_bundle(bundle)
        generate_candidates(bundle)
        if wbs not in {task.wbs for task in bundle.tasks}:
            raise ValueError("找不到需要调整的任务")
        if not 0 <= actual_percent <= 100:
            raise ValueError("实际完成率必须在 0 到 100 之间")
        if actual_percent < 100 and remaining_duration <= 0:
            raise ValueError("未完成任务的剩余工期必须大于 0")
        if not reason.strip() or not action_type.strip() or not note.strip() or not actor.strip():
            raise ValueError("偏差原因、干预措施、调整说明和操作人不能为空")
        remaining_duration = 0 if actual_percent == 100 else remaining_duration
        target_run_id = run_id
        if metadata["state"] != "草案":
            target_run_id = "RUN-" + uuid.uuid4().hex[:10].upper()
            version = self.store.next_version(metadata["project_id"])
            run_dir = self.generated_dir / target_run_id
            run_dir.mkdir(parents=True, exist_ok=False)
            input_path = run_dir / "input.xlsx"
            input_path.write_bytes(content)
            self.store.create_adjustment_run(run_id, target_run_id, version, input_path, actor.strip())
        self.store.save_manual_adjustment(
            target_run_id, wbs, actual_percent, remaining_duration,
            reason.strip(), action_type.strip(), note.strip(), actor.strip(),
        )
        result = self._recalculate(target_run_id)
        self._write_outputs(result, self.generated_dir / target_run_id)
        return self.get_result(target_run_id)

    def transition(self, run_id: str, action: str, actor: str, comment: str) -> dict:
        self.store.transition(run_id, action, actor, comment)
        payload = self.get_result(run_id)
        result = self._recalculate(run_id)
        self._write_outputs(result, self.generated_dir / run_id)
        return payload

    def chat(self, run_id: str, message: str) -> str:
        result = self.get_result(run_id)
        text = message.strip()
        summary = result["summary"]
        if any(keyword in text for keyword in ["关键", "线路"]):
            path = " → ".join(summary.get("critical_path") or [])
            return f"当前基线关键任务为：{path or '无可识别关键线路'}。滚动预测中的关键性请结合“进度偏差”页查看。"
        if any(keyword in text for keyword in ["延期", "偏差", "风险"]):
            alerts = [row for row in result["forecast"] if row["alert"]]
            if not alerts:
                return "当前没有达到组合预警阈值的任务。建议继续按周更新实际量和剩余工期。"
            details = "；".join(f"{row['wbs']}：{row['alert_reason']}" for row in alerts[:5])
            return f"共 {len(alerts)} 项重点预警。{details}。"
        if any(keyword in text for keyword in ["纠偏", "建议", "追回"]):
            scenarios = result["scenarios"]
            if not scenarios:
                return "当前未生成纠偏情景，因为没有达到阈值的关键任务。"
            details = "；".join(f"{row['scenario_id']} 针对 {row['task_wbs']}，预计追回 {row['recovered_days']} 个工作日" if row["quantified"] else f"{row['scenario_id']} 针对 {row['task_wbs']}，需先补充赶工参数" for row in scenarios)
            return details + "。所有情景需计划工程师复核、项目经理批准后执行。"
        if any(keyword in text for keyword in ["问题", "缺失", "正式"]):
            blocking = [row for row in result["issues"] if row["blocking"]]
            if not blocking:
                return f"当前无阻断问题，运行状态为“{result['state']}”，具备进入复核流程的条件。"
            return f"当前有 {len(blocking)} 项阻断问题：" + "；".join(f"{row['entity_id'] or '项目'} {row['message']}" for row in blocking[:8]) + "。"
        finish = summary.get("forecast_finish") or "无法计算"
        variance = summary.get("variance_contract_days")
        return f"本次纳入 {summary['task_count']} 项任务，预测竣工 {finish}，对合同偏差 {variance if variance is not None else '-'} 个工作日，重点预警 {summary['alert_count']} 项。你可以继续问关键线路、延期风险、纠偏建议或阻断问题。"

    def artifact_path(self, run_id: str, kind: str) -> Path:
        names = {"excel": "进度计划分析.xlsx", "pdf": "进度计划分析报告.pdf", "input": "input.xlsx"}
        if kind not in names:
            raise ValueError("不支持的文件类型")
        path = self.generated_dir / run_id / names[kind]
        if not path.exists():
            raise ValueError("文件不存在")
        return path

    def workbook_preview(self, run_id: str, kind: str, sheet_name: str = "") -> dict:
        path = self.artifact_path(run_id, kind)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("仅支持预览 Excel 文件")
        workbook = load_workbook(path, data_only=False, read_only=True)
        selected = sheet_name if sheet_name in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[selected]
        sheet_names = list(workbook.sheetnames)
        truncated = sheet.max_row > 200 or sheet.max_column > 40
        rows = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 200), max_col=min(sheet.max_column, 40)),
            start=1,
        ):
            rows.append([
                {"cell": f"{get_column_letter(column_number)}{row_number}", "value": self._preview_value(cell.value)}
                for column_number, cell in enumerate(row, start=1)
            ])
        workbook.close()
        return {
            "kind": kind,
            "filename": path.name if kind != "input" else self.store.get_run(run_id)["input_filename"],
            "sheets": sheet_names,
            "sheet": selected,
            "rows": rows,
            "truncated": truncated,
        }

    @staticmethod
    def _preview_value(value):
        if value is None:
            return ""
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return str(value)

    def _recalculate(self, run_id: str, use_stored_eligibility: bool = True):
        metadata = self.store.get_run(run_id)
        content = Path(metadata["input_path"]).read_bytes()
        bundle = read_bundle(content, metadata["input_filename"])
        self._normalize_bundle(bundle)
        candidate_notes = generate_candidates(bundle)
        import_confirmed = bool(metadata.get("import_confirmed"))
        if import_confirmed:
            self._confirm_auto_content(bundle)
        adjustments = self.store.get_manual_adjustments(run_id)
        self._apply_manual_adjustments(bundle, adjustments)
        result = analyze_bundle(bundle, run_id, metadata["version"])
        self._apply_candidate_issues(result, candidate_notes, import_confirmed)
        self._apply_import_issues(result, bundle, import_confirmed)
        if use_stored_eligibility:
            result.formal_eligible = bool(metadata["formal_eligible"])
        result.state = metadata["state"]
        self._append_manual_traceability(result, adjustments)
        return result

    def _write_outputs(self, result, run_dir: Path) -> None:
        export_result(result, run_dir / "进度计划分析.xlsx")
        export_pdf(result, run_dir / "进度计划分析报告.pdf")
        (run_dir / "result.json").write_text(json.dumps(result.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")

    def _payload(self, result, import_confirmed: bool = False, adjustments: list[dict] | None = None) -> dict:
        payload = result.to_dict()
        confirmation_available = any(issue.code in {"AUTO_IMPORT_REVIEW", "AUTO_WBS_UNCONFIRMED"} for issue in result.issues)
        confirmable_codes = {"AUTO_IMPORT_REVIEW", "AUTO_WBS_UNCONFIRMED", "UNCONFIRMED_WBS", "UNCONFIRMED_RELATION"}
        payload["import_confirmed"] = import_confirmed
        payload["requires_import_confirmation"] = confirmation_available and not import_confirmed
        payload["unresolved_blocking_count"] = sum(
            issue.blocking and (not confirmation_available or issue.code not in confirmable_codes)
            for issue in result.issues
        )
        payload["manual_adjustments"] = adjustments or []
        payload["downloads"] = {
            "excel": f"/api/runs/{result.run_id}/files/excel",
            "pdf": f"/api/runs/{result.run_id}/files/pdf",
            "input": f"/api/runs/{result.run_id}/files/input",
        }
        return payload

    @staticmethod
    def _candidate_issue(note):
        from .models import Issue
        return Issue("错误", "AUTO_WBS_UNCONFIRMED", f"自动拆解候选待确认：{note}", blocking=True)

    def _apply_candidate_issues(self, result, candidate_notes: list[str], confirmed: bool = False) -> None:
        for note in candidate_notes:
            if confirmed:
                from .models import Issue
                result.issues.insert(0, Issue("警告", "AUTO_WBS_CONFIRMED", f"已人工确认自动拆解：{note}", blocking=False))
            else:
                result.issues.insert(0, self._candidate_issue(note))
        if candidate_notes and not confirmed:
            result.formal_eligible = False
            result.summary["blocking_count"] += len(candidate_notes)

    @staticmethod
    def _apply_import_issues(result, bundle, confirmed: bool = False) -> None:
        if bundle.import_mode == "标准模板":
            return
        from .models import Issue

        result.issues.insert(0, Issue(
            "警告" if confirmed else "错误",
            "AUTO_IMPORT_CONFIRMED" if confirmed else "AUTO_IMPORT_REVIEW",
            "普通进度表的字段映射、默认值和推断关系已由计划工程师确认。" if confirmed else "普通进度表已自动识别并补充为分析草案；字段映射、默认值和推断关系必须由计划工程师确认。",
            bundle.project.project_id,
            "自动导入",
            not confirmed,
        ))
        for note in reversed(bundle.import_notes):
            result.issues.insert(1, Issue("警告", "AUTO_IMPORT_NOTE", note, field_name="自动导入", blocking=False))
        if not confirmed:
            result.formal_eligible = False
            result.summary["blocking_count"] += 1

    @staticmethod
    def _confirm_auto_content(bundle) -> None:
        for task in bundle.tasks:
            task.confirmed = True
        for relation in bundle.relations:
            relation.confirmed = True

    @staticmethod
    def _apply_manual_adjustments(bundle, adjustments: list[dict]) -> None:
        if not adjustments:
            return
        from .models import ActualProgress, DelayCause, SourceRef

        tasks = {task.wbs: task for task in bundle.tasks}
        actuals = {actual.wbs: actual for actual in bundle.actuals}
        for adjustment in adjustments:
            task = tasks.get(adjustment["wbs"])
            if not task:
                continue
            percent = float(adjustment["actual_percent"])
            existing = actuals.get(task.wbs)
            if percent >= 100:
                task.status = "已完成"
                actual_start = existing.actual_start if existing else bundle.project.data_date
                actual_finish = bundle.project.data_date
                remaining = 0
            elif percent > 0:
                task.status = "进行中"
                actual_start = existing.actual_start if existing and existing.actual_start else bundle.project.data_date
                actual_finish = None
                remaining = int(adjustment["remaining_duration"])
            else:
                task.status = "未开始"
                actual_start = None
                actual_finish = None
                remaining = int(adjustment["remaining_duration"])
            actual = ActualProgress(
                task.wbs,
                task.quantity * percent / 100.0,
                actual_start,
                actual_finish,
                remaining,
                f"人工调整#{adjustment['id']}",
                adjustment["actor"],
            )
            if existing:
                bundle.actuals[bundle.actuals.index(existing)] = actual
            else:
                bundle.actuals.append(actual)
            actuals[task.wbs] = actual
            bundle.delay_causes.append(DelayCause(
                task.wbs, adjustment["reason"], adjustment["note"], "人工判定", "待确认"
            ))
            cell = f"记录#{adjustment['id']}"
            bundle.source_refs.extend([
                SourceRef("人工调整", task.wbs, "实际完成工程量", "人工调整", cell, actual.actual_quantity),
                SourceRef("人工调整", task.wbs, "剩余工期", "人工调整", cell, remaining),
                SourceRef("人工调整", task.wbs, "偏差原因", "人工调整", cell, adjustment["reason"]),
                SourceRef("人工调整", task.wbs, "干预措施", "人工调整", cell, adjustment["action_type"]),
            ])

    @staticmethod
    def _append_manual_traceability(result, adjustments: list[dict]) -> None:
        for adjustment in adjustments:
            inputs = f"人工调整!记录#{adjustment['id']}"
            for field, value in (
                ("偏差原因", adjustment["reason"]),
                ("干预措施", adjustment["action_type"]),
                ("调整说明", adjustment["note"]),
            ):
                result.traceability.append({
                    "output": "人工进度调整",
                    "entity_id": adjustment["wbs"],
                    "field": field,
                    "value": value,
                    "input_refs": inputs,
                    "rule_version": result.rule_version,
                    "rule": "计划工程师人工录入，重算滚动预测，不修改批准基线",
                    "calculated_at": result.calculated_at,
                    "input_hash": result.input_hash,
                    "manual_override": "是",
                })

    @staticmethod
    def _normalize_bundle(bundle) -> None:
        bundle.project.project_id = str(bundle.project.project_id or "").strip()
        bundle.project.project_name = str(bundle.project.project_name or "").strip()
        bundle.project.calendar_id = str(bundle.project.calendar_id or "").strip()
        bundle.calendar.calendar_id = str(bundle.calendar.calendar_id or "").strip()
        text_fields = {
            "tasks": ("wbs", "name", "discipline", "section", "station", "workfront", "unit", "status"),
            "relations": ("predecessor", "successor", "relation_type"),
            "resources": ("wbs", "resource_type", "source", "applicable_condition"),
            "crash_params": ("wbs", "overtime_condition", "approval_required"),
            "constraints": ("constraint_id", "wbs", "category", "owner", "status"),
            "actuals": ("wbs", "evidence_ref", "reporter"),
            "delay_causes": ("wbs", "category", "note", "responsibility", "controllable"),
            "period_plans": ("wbs",),
            "baselines": ("wbs",),
            "milestones": ("milestone_id", "name", "linked_wbs"),
        }
        for collection_name, fields in text_fields.items():
            for item in getattr(bundle, collection_name):
                for field in fields:
                    value = getattr(item, field)
                    setattr(item, field, "" if value is None else str(value).strip())
        for task in bundle.tasks:
            task.quantity = float(task.quantity or 0)
            task.duration_override = int(task.duration_override) if task.duration_override not in (None, "") else None
        for relation in bundle.relations:
            relation.relation_type = relation.relation_type.upper()
            relation.lag = int(relation.lag or 0)
        for actual in bundle.actuals:
            actual.actual_quantity = float(actual.actual_quantity or 0)
            actual.remaining_duration = int(actual.remaining_duration) if actual.remaining_duration not in (None, "") else None
        for source in bundle.source_refs:
            source.entity_type = str(source.entity_type or "").strip()
            source.entity_id = str(source.entity_id or "").strip()
            source.field_name = str(source.field_name or "").strip()

    def _record_import_error(self, filename: str, stage: str, exc: Exception) -> None:
        log_path = self.data_dir / "import_errors.log"
        timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
        entry = (
            f"[{timestamp}] file={filename!r} stage={stage} error={type(exc).__name__}: {exc}\n"
            f"{traceback.format_exc()}\n"
        )
        with log_path.open("a", encoding="utf-8") as stream:
            stream.write(entry)

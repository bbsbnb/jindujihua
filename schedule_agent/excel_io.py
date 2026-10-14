from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    ActualProgress,
    BaselineRecord,
    CalendarSpec,
    CrashParam,
    DelayCause,
    ExternalConstraint,
    InputBundle,
    Milestone,
    PeriodPlan,
    Project,
    Relation,
    Resource,
    SourceRef,
    Task,
)


SHEETS = {
    "项目参数": ["参数", "值"],
    "任务与工程量": ["WBS编码", "任务名称", "专业", "区段", "桩号范围", "工作面", "计量单位", "总工程量", "任务状态", "工期覆盖", "是否里程碑", "是否确认"],
    "逻辑关系": ["前置任务", "后续任务", "关系类型", "时滞工作日", "是否确认"],
    "资源与定额": ["WBS编码", "资源类型", "当前数量", "硬上限", "可增调量", "单位日产能", "定额来源", "适用条件", "是否确认"],
    "赶工参数": ["WBS编码", "最大资源投入", "最短可行工期", "资源效率系数", "加班条件", "审批要求", "是否确认"],
    "外部约束": ["约束编号", "WBS编码", "约束类别", "最早可开始日期", "责任方", "状态", "是否确认"],
    "工作日历": ["日历编号", "每周工作日", "节假日", "额外工作日"],
    "实际进度": ["WBS编码", "实际完成工程量", "实际开始日期", "实际完成日期", "剩余工期", "证据来源编号", "填报人"],
    "偏差原因": ["WBS编码", "原因分类", "补充说明", "责任类别", "是否可控"],
    "时段计划量": ["WBS编码", "期间截止日", "计划工程量"],
    "基线": ["WBS编码", "合同开始", "合同完成", "当前批准开始", "当前批准完成"],
    "里程碑": ["里程碑编号", "里程碑名称", "关联WBS", "合同日期", "当前批准日期"],
}

PARAMETER_KEYS = ["项目编号", "项目名称", "项目开工日", "合同竣工日", "当前批准竣工日", "数据日期", "日历编号", "编制人", "项目经理"]
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


class WorkbookInputError(ValueError):
    """A user-facing workbook format or structure error."""


def create_template(path: Path, sample: bool = False) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in SHEETS.items():
        ws = workbook.create_sheet(sheet_name)
        ws.append(headers)
        style_sheet(ws, headers)

    project_values = {
        "项目编号": "SZDL-DEMO-001" if sample else "",
        "项目名称": "滨河路综合改造工程（示例）" if sample else "",
        "项目开工日": date(2026, 10, 12) if sample else "",
        "合同竣工日": date(2026, 12, 31) if sample else "",
        "当前批准竣工日": date(2026, 12, 31) if sample else "",
        "数据日期": date(2026, 10, 23) if sample else "",
        "日历编号": "CAL-01",
        "编制人": "张计划" if sample else "",
        "项目经理": "李经理" if sample else "",
    }
    ws = workbook["项目参数"]
    for key in PARAMETER_KEYS:
        ws.append([key, project_values[key]])

    workbook["工作日历"].append(["CAL-01", "1,2,3,4,5", "2026-10-01,2026-10-02", ""])
    if sample:
        _populate_sample(workbook)

    for ws in workbook.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            width = min(32, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd"
    workbook.save(path)


def _populate_sample(workbook: Workbook) -> None:
    tasks = [
        ["1.1", "交通导改及围挡", "临时工程", "A区", "K0+000-K0+500", "道路主线", "m", 500, "已完成", "", "否", "是"],
        ["1.2", "雨水管沟开挖", "排水", "A区", "K0+000-K0+500", "东侧管沟", "m", 500, "进行中", "", "否", "是"],
        ["1.3", "雨水管安装回填", "排水", "A区", "K0+000-K0+500", "东侧管沟", "m", 500, "未开始", "", "否", "是"],
        ["2.1", "路床整形", "道路", "A区", "K0+000-K0+500", "道路主线", "m2", 6000, "未开始", "", "否", "是"],
        ["2.2", "水稳基层施工", "道路", "A区", "K0+000-K0+500", "道路主线", "m2", 6000, "未开始", "", "否", "是"],
        ["2.3", "沥青面层施工", "道路", "A区", "K0+000-K0+500", "道路主线", "m2", 6000, "未开始", "", "否", "是"],
        ["3.1", "照明与附属设施", "照明", "A区", "K0+000-K0+500", "道路两侧", "套", 40, "未开始", "", "否", "是"],
        ["M1", "项目完工里程碑", "项目", "A区", "全线", "全线", "项", 0, "未开始", 1, "是", "是"],
    ]
    for row in tasks:
        workbook["任务与工程量"].append(row)
    relations = [
        ["1.1", "1.2", "FS", 0, "是"],
        ["1.2", "1.3", "SS", 5, "是"],
        ["1.3", "2.1", "FS", 0, "是"],
        ["2.1", "2.2", "FS", 0, "是"],
        ["2.2", "2.3", "FS", 0, "是"],
        ["2.2", "3.1", "SS", 3, "是"],
        ["2.3", "M1", "FF", 0, "是"],
        ["3.1", "M1", "FS", 0, "是"],
    ]
    for row in relations:
        workbook["逻辑关系"].append(row)
    resource_rows = [
        ["1.1", "综合班组", 1, 2, 1, 100, "企业定额", "围挡材料到场", "是"],
        ["1.2", "管沟班组", 1, 2, 1, 35, "项目实测", "正常土质", "是"],
        ["1.3", "管道班组", 1, 2, 1, 40, "企业定额", "沟槽验收合格", "是"],
        ["2.1", "路基班组", 1, 2, 1, 600, "企业定额", "工作面移交", "是"],
        ["2.2", "基层班组", 1, 2, 1, 500, "企业定额", "拌合站稳定供应", "是"],
        ["2.3", "摊铺机组", 1, 1, 0, 750, "企业定额", "连续摊铺", "是"],
        ["3.1", "安装班组", 1, 2, 1, 4, "企业定额", "基础具备条件", "是"],
    ]
    for row in resource_rows:
        workbook["资源与定额"].append(row)
    for row in [
        ["1.2", 2, 8, 0.75, "可增加夜班", "项目经理批准", "是"],
        ["1.3", 2, 7, 0.70, "工作面满足", "项目经理批准", "是"],
        ["2.2", 2, 8, 0.65, "材料连续供应", "项目经理批准", "是"],
    ]:
        workbook["赶工参数"].append(row)
    for row in [
        ["C01", "1.1", "交通导改", date(2026, 10, 12), "交管协调组", "已完成", "是"],
        ["C02", "2.3", "材料审批", date(2026, 11, 30), "物资部", "进行中", "是"],
    ]:
        workbook["外部约束"].append(row)
    for row in [
        ["1.1", 500, date(2026, 10, 12), date(2026, 10, 16), 0, "周报-2026-W42", "张计划"],
        ["1.2", 120, date(2026, 10, 19), "", 14, "周报-2026-W43", "张计划"],
    ]:
        workbook["实际进度"].append(row)
    workbook["偏差原因"].append(["1.2", "地下管线影响", "既有通信管线位置与图纸不符", "外部", "否"])
    for row in [
        ["1.1", date(2026, 10, 16), 500],
        ["1.2", date(2026, 10, 23), 175],
    ]:
        workbook["时段计划量"].append(row)
    workbook["里程碑"].append(["MS-01", "道路开放交通", "M1", date(2026, 12, 31), date(2026, 12, 31)])


def read_bundle(content: bytes, filename: str) -> InputBundle:
    _validate_workbook_container(content, filename)
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise WorkbookInputError(
            "Excel 工作簿结构无效或已经损坏。请用 Excel/WPS 打开原文件，取消密码保护后另存为“Excel 工作簿 (*.xlsx)”，再重新导入。"
        ) from exc
    missing = [name for name in SHEETS if name not in workbook.sheetnames]
    if missing:
        from .generic_import import read_generic_workbook

        return read_generic_workbook(workbook, content, filename)
    refs: list[SourceRef] = []

    params = {}
    for row in workbook["项目参数"].iter_rows(min_row=2):
        key, value = row[0].value, row[1].value
        if key:
            params[str(key).strip()] = value
            refs.append(SourceRef("项目", str(params.get("项目编号", "")), str(key), "项目参数", row[1].coordinate, value))
    required_params = ["项目编号", "项目名称", "项目开工日", "合同竣工日", "数据日期", "日历编号"]
    absent = [key for key in required_params if params.get(key) in (None, "")]
    if absent:
        raise ValueError(f"项目参数缺少: {', '.join(absent)}")
    project = Project(
        project_id=str(params["项目编号"]),
        project_name=str(params["项目名称"]),
        start_date=_date(params["项目开工日"]),
        contract_finish=_date(params["合同竣工日"]),
        current_finish=_date_or_none(params.get("当前批准竣工日")),
        data_date=_date(params["数据日期"]),
        calendar_id=str(params["日历编号"]),
        preparer=str(params.get("编制人") or ""),
        manager=str(params.get("项目经理") or ""),
    )

    calendars = _records(workbook["工作日历"], refs, "日历")
    calendar_row = next((item for item in calendars if str(item["日历编号"]) == project.calendar_id), None)
    if not calendar_row:
        raise ValueError(f"找不到项目日历: {project.calendar_id}")
    calendar = CalendarSpec(
        calendar_id=project.calendar_id,
        work_weekdays={int(value.strip()) - 1 for value in str(calendar_row["每周工作日"]).split(",") if value.strip()},
        holidays=set(_date_list(calendar_row.get("节假日"))),
        extra_workdays=set(_date_list(calendar_row.get("额外工作日"))),
    )

    tasks = [Task(
        wbs=str(row["WBS编码"] or "").strip(), name=str(row["任务名称"] or "").strip(), discipline=str(row["专业"] or ""),
        section=str(row["区段"] or ""), station=str(row["桩号范围"] or ""), workfront=str(row["工作面"] or ""),
        unit=str(row["计量单位"] or ""), quantity=_float(row["总工程量"]), status=str(row["任务状态"] or "未开始"),
        duration_override=_int_or_none(row.get("工期覆盖")), is_milestone=_bool(row.get("是否里程碑")), confirmed=_bool(row.get("是否确认")),
    ) for row in _records(workbook["任务与工程量"], refs, "任务") if any(value not in (None, "") for value in row.values())]
    relations = [Relation(str(row["前置任务"]), str(row["后续任务"]), str(row["关系类型"]).upper(), int(row["时滞工作日"] or 0), _bool(row["是否确认"])) for row in _records(workbook["逻辑关系"], refs, "逻辑") if row["前置任务"]]
    resources = [Resource(str(row["WBS编码"]), str(row["资源类型"]), _float(row["当前数量"]), _float(row["硬上限"]), _float(row["可增调量"]), _float(row["单位日产能"]), str(row["定额来源"] or ""), str(row["适用条件"] or ""), _bool(row["是否确认"])) for row in _records(workbook["资源与定额"], refs, "资源") if row["WBS编码"]]
    crash_params = [CrashParam(str(row["WBS编码"]), _float(row["最大资源投入"]), int(row["最短可行工期"]), _float(row["资源效率系数"]), str(row["加班条件"] or ""), str(row["审批要求"] or ""), _bool(row["是否确认"])) for row in _records(workbook["赶工参数"], refs, "赶工") if row["WBS编码"]]
    constraints = [ExternalConstraint(str(row["约束编号"]), str(row["WBS编码"]), str(row["约束类别"]), _date(row["最早可开始日期"]), str(row["责任方"] or ""), str(row["状态"] or ""), _bool(row["是否确认"])) for row in _records(workbook["外部约束"], refs, "约束") if row["约束编号"]]
    actuals = [ActualProgress(str(row["WBS编码"]), _float(row["实际完成工程量"]), _date_or_none(row["实际开始日期"]), _date_or_none(row["实际完成日期"]), _int_or_none(row["剩余工期"]), str(row["证据来源编号"] or ""), str(row["填报人"] or "")) for row in _records(workbook["实际进度"], refs, "实际") if row["WBS编码"]]
    delay_causes = [DelayCause(str(row["WBS编码"]), str(row["原因分类"] or ""), str(row["补充说明"] or ""), str(row["责任类别"] or ""), str(row["是否可控"] or "")) for row in _records(workbook["偏差原因"], refs, "原因") if row["WBS编码"]]
    period_plans = [PeriodPlan(str(row["WBS编码"]), _date(row["期间截止日"]), _float(row["计划工程量"])) for row in _records(workbook["时段计划量"], refs, "时段计划") if row["WBS编码"]]
    baselines = [BaselineRecord(str(row["WBS编码"]), _date_or_none(row["合同开始"]), _date_or_none(row["合同完成"]), _date_or_none(row["当前批准开始"]), _date_or_none(row["当前批准完成"])) for row in _records(workbook["基线"], refs, "基线") if row["WBS编码"]]
    milestones = [Milestone(str(row["里程碑编号"]), str(row["里程碑名称"]), str(row["关联WBS"]), _date(row["合同日期"]), _date_or_none(row["当前批准日期"])) for row in _records(workbook["里程碑"], refs, "里程碑") if row["里程碑编号"]]
    return InputBundle(project, calendar, tasks, relations, resources, crash_params, constraints, actuals, delay_causes, period_plans, baselines, milestones, refs, hashlib.sha256(content).hexdigest(), filename)


def _validate_workbook_container(content: bytes, filename: str) -> None:
    if not filename.lower().endswith(".xlsx"):
        raise WorkbookInputError("仅支持 .xlsx 格式。请将文件另存为“Excel 工作簿 (*.xlsx)”后重新导入。")
    if not content:
        raise WorkbookInputError("所选文件为空，请重新选择有效的 Excel 工作簿。")
    if content.startswith(OLE_SIGNATURE):
        raise WorkbookInputError(
            "该文件实际是旧版 .xls 或受密码保护的 Office 文件，并非可解析的 .xlsx。请取消密码保护，并另存为“Excel 工作簿 (*.xlsx)”。"
        )
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise WorkbookInputError(
            "该文件不是有效的 .xlsx 工作簿，可能只是修改了文件后缀。请用 Excel/WPS 打开后另存为“Excel 工作簿 (*.xlsx)”。"
        )
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = set(archive.namelist())
            required = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"}
            if not required.issubset(names):
                raise WorkbookInputError(
                    "文件中缺少 Excel 工作簿主体，可能是损坏文件或其他 Office 文件。请重新下载标准模板，或另存为新的 .xlsx 后再导入。"
                )
            if archive.testzip() is not None:
                raise WorkbookInputError("Excel 文件压缩结构已损坏，请从原始文件重新另存为 .xlsx 后导入。")
    except zipfile.BadZipFile as exc:
        raise WorkbookInputError("Excel 文件压缩结构已损坏，请重新获取文件后导入。") from exc


def _records(ws, refs: list[SourceRef], entity_type: str) -> list[dict[str, Any]]:
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    rows = []
    for excel_row in ws.iter_rows(min_row=2):
        record = {headers[index]: cell.value for index, cell in enumerate(excel_row) if index < len(headers)}
        if not any(value not in (None, "") for value in record.values()):
            continue
        entity_id = str(next((record.get(key) for key in ["WBS编码", "约束编号", "里程碑编号", "日历编号"] if record.get(key)), ws.title))
        for index, cell in enumerate(excel_row):
            if index < len(headers) and cell.value not in (None, ""):
                refs.append(SourceRef(entity_type, entity_id, headers[index], ws.title, cell.coordinate, cell.value))
        rows.append(record)
    return rows


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date.fromisoformat(value.strip())
    raise ValueError(f"无法识别日期: {value}")


def _date_or_none(value: Any) -> date | None:
    return None if value in (None, "") else _date(value)


def _date_list(value: Any) -> list[date]:
    if value in (None, ""):
        return []
    return [_date(item.strip()) for item in str(value).replace("；", ",").split(",") if item.strip()]


def _float(value: Any) -> float:
    return float(value or 0)


def _int_or_none(value: Any) -> int | None:
    return None if value in (None, "") else int(value)


def _bool(value: Any) -> bool:
    return str(value or "").strip().lower() in {"是", "true", "1", "yes", "y"}


def style_sheet(ws, headers) -> None:
    fill = PatternFill("solid", fgColor="17324D")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26


def export_result(result, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_summary(workbook, result)
    _write_rows(workbook, "进度计划表", ["WBS", "任务", "工期", "最早开始", "最早完成", "最迟开始", "最迟完成", "总时差", "自由时差", "关键", "区段", "工作面"], [[r.wbs, r.name, r.duration, r.early_start, r.early_finish, r.late_start, r.late_finish, r.total_float, r.free_float, "是" if r.critical else "否", r.section, r.workfront] for r in result.schedule])
    _write_gantt(workbook, result)
    _write_rows(workbook, "关键线路", ["WBS", "任务", "开始", "完成", "总时差"], [[r.wbs, r.name, r.early_start, r.early_finish, r.total_float] for r in result.schedule if r.critical])
    _write_rows(workbook, "进度偏差", ["WBS", "状态", "计划完成率%", "实际完成率%", "偏差百分点", "当前批准完成", "合同基线完成", "预测完成", "对当前偏差工作日", "对合同偏差工作日", "剩余工期", "当前关键", "预警", "预警原因", "效率推算剩余工期"], [[r.wbs, r.status, r.planned_percent, r.actual_percent, r.variance_pp, r.baseline_finish, r.contract_baseline_finish, r.forecast_finish, r.finish_variance_days, r.contract_variance_days, r.remaining_duration, "是" if r.current_critical else "否", "是" if r.alert else "否", r.alert_reason, r.efficiency_remaining] for r in result.forecast])
    _write_rows(workbook, "里程碑", ["编号", "名称", "关联WBS", "合同日期", "当前批准日期", "预测日期", "对当前偏差", "对合同偏差", "预警"], [[m["milestone_id"], m["name"], m["linked_wbs"], m["contract_date"], m["current_date"], m["forecast_date"], m["variance_current_days"], m["variance_contract_days"], "是" if m["alert"] else "否"] for m in result.milestones])
    _write_rows(workbook, "资源冲突", ["类型", "资源/工作面", "日期", "需求", "能力", "任务", "说明"], [[r.conflict_type, r.key, r.conflict_date, r.demand, r.capacity, ",".join(r.tasks), r.message] for r in result.conflicts])
    _write_rows(workbook, "纠偏情景", ["情景", "措施", "WBS", "任务", "调整前工期", "调整后工期", "资源增量", "预计追回工作日", "里程碑影响", "资源冲突", "风险", "审批", "已量化"], [[r.scenario_id, r.action_type, r.task_wbs, r.task_name, r.before_duration, r.after_duration, r.resource_delta, r.recovered_days, r.milestone_impact, r.resource_conflict, r.risk, r.approval, "是" if r.quantified else "否"] for r in result.scenarios])
    _write_rows(workbook, "问题清单", ["级别", "编码", "说明", "对象", "字段", "阻断正式输出"], [[r.severity, r.code, r.message, r.entity_id, r.field_name, "是" if r.blocking else "否"] for r in result.issues])
    _write_rows(workbook, "版本记录", ["运行编号", "版本", "状态", "正式输出资格", "规则版本", "计算时间", "输入文件", "输入SHA256"], [[result.run_id, result.version, result.state, "是" if result.formal_eligible else "否", result.rule_version, result.calculated_at, result.input_filename, result.input_hash]])
    _write_rows(workbook, "追溯记录", ["交付物", "对象", "字段", "结果值", "输入引用", "规则版本", "处理规则", "计算时间", "输入SHA256", "人工覆盖"], [[r["output"], r["entity_id"], r["field"], r["value"], r["input_refs"], r["rule_version"], r["rule"], r["calculated_at"], r["input_hash"], r["manual_override"]] for r in result.traceability])
    workbook.save(path)


def _write_summary(workbook: Workbook, result) -> None:
    ws = workbook.create_sheet("管理摘要")
    ws.append([result.project.project_name, "进度计划分析"])
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(size=18, bold=True, color="17324D")
    ws["A1"].alignment = Alignment(vertical="center")
    items = [
        ("运行编号", result.run_id), ("状态", result.state), ("数据日期", result.project.data_date),
        ("合同竣工", result.summary["contract_finish"]), ("预测竣工", result.summary["forecast_finish"]),
        ("对合同偏差（工作日）", result.summary["variance_contract_days"]), ("任务数", result.summary["task_count"]),
        ("当前关键任务", result.summary["critical_count"]), ("重点预警", result.summary["alert_count"]),
        ("阻断问题", result.summary["blocking_count"]), ("正式输出资格", "具备" if result.formal_eligible else "不具备"),
        ("规则版本", result.rule_version),
    ]
    for index, (label, value) in enumerate(items, start=3):
        ws.cell(index, 1, label).font = Font(bold=True, color="456173")
        ws.cell(index, 2, value)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32


def _write_rows(workbook: Workbook, name: str, headers: list[str], rows: list[list]) -> None:
    ws = workbook.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[get_column_letter(column[0].column)].width = width
        for cell in column:
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
    if name in {"进度偏差", "问题清单"} and ws.max_row > 1:
        formula_col = "M" if name == "进度偏差" else "F"
        ws.conditional_formatting.add(f"A2:{get_column_letter(ws.max_column)}{ws.max_row}", FormulaRule(formula=[f'${formula_col}2="是"'], fill=PatternFill("solid", fgColor="FDE7E7")))


def _write_gantt(workbook: Workbook, result) -> None:
    ws = workbook.create_sheet("甘特图")
    ws.append(["WBS", "任务", "起始偏移", "工期", "关键"])
    if not result.schedule:
        style_sheet(ws, ws[1])
        return
    origin = min(row.early_start for row in result.schedule)
    for row in result.schedule:
        ws.append([row.wbs, row.name, (row.early_start - origin).days, row.duration, "是" if row.critical else "否"])
    style_sheet(ws, ws[1])
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "进度横道图（工作日工期）"
    chart.y_axis.title = "任务"
    chart.x_axis.title = f"自 {origin.isoformat()} 起的日历偏移"
    data = Reference(ws, min_col=3, max_col=4, min_row=1, max_row=ws.max_row)
    categories = Reference(ws, min_col=2, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = max(7, min(16, len(result.schedule) * 0.55))
    chart.width = 24
    chart.series[0].graphicalProperties.noFill = True
    chart.series[0].graphicalProperties.line.noFill = True
    chart.series[1].graphicalProperties.solidFill = "2D7D7B"
    ws.add_chart(chart, "G2")
    ws.column_dimensions["B"].width = 30

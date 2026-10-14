from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook

from schedule_agent.excel_io import read_bundle
from schedule_agent.service import ScheduleService


class EndToEndTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = ScheduleService(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def test_sample_analysis_and_exports_are_consistent(self):
        result = self.service.analyze_sample()
        self.assertTrue(result["formal_eligible"])
        self.assertEqual(result["summary"]["baseline_finish"], "2026-12-23")
        self.assertEqual(result["summary"]["forecast_finish"], "2026-12-30")
        self.assertEqual(result["summary"]["variance_contract_days"], -1)
        self.assertIn("1.2", result["summary"]["critical_path"])
        quantified = next(row for row in result["scenarios"] if row["quantified"])
        self.assertGreaterEqual(quantified["recovered_days"], 0)
        self.assertIn("项目预测竣工由", quantified["milestone_impact"])
        recoveries = [row["recovered_days"] for row in result["scenarios"] if row["recovered_days"] is not None]
        self.assertEqual(recoveries, sorted(recoveries, reverse=True))
        self.assertTrue(any("实际进度!" in row["input_refs"] for row in result["traceability"] if row["field"] == "实际完成率" and row["entity_id"] == "1.2"))

        excel_path = self.service.artifact_path(result["run_id"], "excel")
        workbook = load_workbook(excel_path, data_only=True)
        self.assertIn("进度偏差", workbook.sheetnames)
        self.assertIn("追溯记录", workbook.sheetnames)
        summary = {workbook["管理摘要"].cell(row, 1).value: workbook["管理摘要"].cell(row, 2).value for row in range(3, 15)}
        self.assertEqual(summary["预测竣工"].date().isoformat(), result["summary"]["forecast_finish"])

        pdf_path = self.service.artifact_path(result["run_id"], "pdf")
        with pdfplumber.open(pdf_path) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        self.assertIn("2026-12-30", text)
        self.assertIn(result["run_id"], text)
        self.assertIn("关键线路", text)

    def test_approval_workflow_and_audit(self):
        result = self.service.analyze_sample()
        run_id = result["run_id"]
        reviewed = self.service.transition(run_id, "review", "张计划", "复核通过")
        self.assertEqual(reviewed["state"], "已复核")
        approved = self.service.transition(run_id, "approve", "李经理", "批准")
        self.assertEqual(approved["state"], "已批准")
        self.assertEqual([row["to_state"] for row in approved["audit"]], ["草案", "已复核", "已批准"])
        with self.assertRaisesRegex(ValueError, "不允许"):
            self.service.transition(run_id, "approve", "李经理", "重复批准")

    def test_delete_plan_archives_all_versions_and_preserves_audit(self):
        first = self.service.analyze_sample()
        second = self.service.analyze_sample()
        result = self.service.archive_plan(second["run_id"], "张计划")
        self.assertEqual(result["archived_runs"], 2)
        self.assertEqual(self.service.store.list_runs(), [])
        self.assertIsNotNone(self.service.store.get_run(first["run_id"]))
        self.assertEqual(self.service.store.audit(first["run_id"])[-1]["action"], "删除计划")

    def test_unconfirmed_rate_blocks_formal_output(self):
        workbook = load_workbook(self.service.sample_path)
        workbook["资源与定额"]["G2"] = "行业参考"
        workbook["资源与定额"]["I2"] = "否"
        broken_path = self.root / "broken.xlsx"
        workbook.save(broken_path)
        result = self.service.analyze(broken_path.read_bytes(), broken_path.name)
        self.assertFalse(result["formal_eligible"])
        self.assertTrue(any(row["code"] == "UNCONFIRMED_RATE" for row in result["issues"]))
        with self.assertRaisesRegex(ValueError, "阻断"):
            self.service.transition(result["run_id"], "review", "张计划", "")

    def test_invalid_workbook_errors_are_actionable(self):
        with self.assertRaisesRegex(ValueError, "不是有效的 .xlsx"):
            self.service.analyze(b"not an excel workbook", "renamed.xlsx")
        with self.assertRaisesRegex(ValueError, "旧版 .xls 或受密码保护"):
            self.service.analyze(bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy", "legacy.xlsx")

    def test_zip_without_workbook_part_is_rejected(self):
        import io
        import zipfile

        payload = io.BytesIO()
        with zipfile.ZipFile(payload, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("_rels/.rels", "<Relationships/>")
        with self.assertRaisesRegex(ValueError, "缺少 Excel 工作簿主体"):
            self.service.analyze(payload.getvalue(), "broken.xlsx")

    def test_generic_schedule_is_auto_mapped_to_draft(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "施工总进度"
        sheet.append(["工程名称", "城东路提升工程"])
        sheet.append(["数据日期", "2026-10-20"])
        sheet.append(["任务编码", "工作内容", "施工段", "计划开始", "计划完成", "完成率", "紧前工作", "实际开始", "剩余工期"])
        sheet.append(["A", "交通导改", "一标段", "2026-10-12", "2026-10-16", "100%", "", "2026-10-12", 0])
        sheet.append(["B", "雨水管施工", "一标段", "2026-10-19", "2026-10-30", "30%", "A FS", "2026-10-19", 7])
        sheet.append(["C", "路床整形", "一标段", "2026-10-31", "2026-11-06", 0, "BFS+1", "", 5])
        path = self.root / "普通施工进度表.xlsx"
        workbook.save(path)

        result = self.service.analyze(path.read_bytes(), path.name)
        self.assertEqual(result["summary"]["import_mode"], "普通进度表自动映射")
        self.assertFalse(result["formal_eligible"])
        self.assertEqual(result["summary"]["task_count"], 3)
        self.assertTrue(result["schedule"])
        self.assertEqual(result["forecast"][0]["actual_percent"], 100.0)
        self.assertTrue(any(row["code"] == "AUTO_IMPORT_REVIEW" for row in result["issues"]))
        self.assertTrue(any("字段映射" in row["message"] for row in result["issues"]))
        self.assertTrue(any("施工总进度!" in row["input_refs"] for row in result["traceability"]))
        self.assertTrue(result["requires_import_confirmation"])
        confirmed = self.service.confirm_import(result["run_id"], "张计划", True)
        self.assertTrue(confirmed["formal_eligible"])
        self.assertFalse(confirmed["requires_import_confirmation"])
        reviewed = self.service.transition(result["run_id"], "review", "张计划", "自动导入内容已核对")
        self.assertEqual(reviewed["state"], "已复核")
        self.assertTrue(any(row["action"] == "确认自动导入" for row in reviewed["audit"]))

    def test_generic_schedule_without_task_table_has_clear_error(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["说明", "这不是进度任务表"])
        path = self.root / "无任务表.xlsx"
        workbook.save(path)
        with self.assertRaisesRegex(ValueError, "至少需要“任务名称”"):
            self.service.analyze(path.read_bytes(), path.name)

    def test_generic_schedule_normalizes_mixed_identifier_types(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["任务编码", "任务名称", "计划开始", "计划完成"])
        sheet.append([1, "数字编码任务", "2026-10-12", "2026-10-13"])
        sheet.append(["2.1", "文本编码任务", "2026-10-14", "2026-10-15"])
        path = self.root / "混合编码.xlsx"
        workbook.save(path)
        result = self.service.analyze(path.read_bytes(), path.name)
        self.assertEqual([row["wbs"] for row in result["schedule"]], ["1", "2.1"])
        reopened = self.service.get_result(result["run_id"])
        self.assertEqual([row["wbs"] for row in reopened["schedule"]], ["1", "2.1"])

    def test_out_of_sequence_keeps_incomplete_relationship(self):
        workbook = load_workbook(self.service.sample_path)
        task_ws = workbook["任务与工程量"]
        # 1.3 starts despite its predecessor 1.2 remaining incomplete.
        for row in range(2, task_ws.max_row + 1):
            if task_ws.cell(row, 1).value == "1.3":
                task_ws.cell(row, 9).value = "进行中"
        workbook["实际进度"].append(["1.3", 20, "2026-10-22", "", 12, "周报-乱序", "张计划"])
        path = self.root / "out-of-sequence.xlsx"
        workbook.save(path)
        result = self.service.analyze(path.read_bytes(), path.name)
        forecast = {row["wbs"]: row for row in result["forecast"]}
        self.assertGreaterEqual(forecast["1.3"]["forecast_finish"], forecast["1.2"]["forecast_start"])


if __name__ == "__main__":
    unittest.main()

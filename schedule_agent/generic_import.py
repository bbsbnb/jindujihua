from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl.utils.datetime import from_excel

from .calendar import WorkCalendar
from .models import (
    ActualProgress,
    BaselineRecord,
    CalendarSpec,
    InputBundle,
    Milestone,
    Project,
    Relation,
    SourceRef,
    Task,
)


ALIASES = {
    "wbs": ["wbs", "wbs编码", "任务编码", "作业代码", "工作代码", "编码", "序号", "id"],
    "name": ["任务名称", "工作名称", "作业名称", "工序名称", "分项工程", "施工内容", "工作内容", "项目名称", "任务"],
    "discipline": ["专业", "专业类别", "工程专业"],
    "section": ["区段", "施工区段", "标段", "施工段", "区域"],
    "station": ["桩号范围", "桩号", "起止桩号"],
    "workfront": ["工作面", "作业面", "施工面"],
    "unit": ["单位", "计量单位"],
    "quantity": ["总工程量", "工程量", "计划工程量", "合同工程量"],
    "actual_quantity": ["实际完成工程量", "已完成工程量", "完成工程量", "累计完成量"],
    "plan_start": ["计划开始日期", "计划开始时间", "计划开工日期", "计划开工", "计划开始", "开始日期"],
    "plan_finish": ["计划完成日期", "计划结束日期", "计划完工日期", "计划完成时间", "计划完成", "计划结束", "计划完工", "结束日期", "完成日期"],
    "duration": ["计划工期", "持续时间", "工期天数", "工期", "duration"],
    "predecessor": ["前置任务", "紧前工作", "紧前任务", "前置关系", "逻辑关系", "前置"],
    "progress": ["实际完成率", "累计完成率", "完成百分比", "完成比例", "完成率", "进度百分比", "实际进度", "进度"],
    "status": ["任务状态", "进度状态", "状态"],
    "actual_start": ["实际开始日期", "实际开工日期", "实际开始时间", "实际开工", "实际开始"],
    "actual_finish": ["实际完成日期", "实际完工日期", "实际结束日期", "实际完成时间", "实际完工", "实际完成", "实际结束"],
    "remaining": ["剩余工期", "尚需工期", "剩余持续时间"],
}

METADATA_ALIASES = {
    "project_name": ["项目名称", "工程名称"],
    "project_id": ["项目编号", "工程编号", "项目代码"],
    "data_date": ["数据日期", "状态日期", "更新日期", "统计日期", "截止日期"],
    "contract_finish": ["合同竣工日", "合同完工日", "合同完成日期", "计划总工期完成日"],
}


class GenericImportError(ValueError):
    pass


def read_generic_workbook(workbook, content: bytes, filename: str) -> InputBundle:
    candidate = _find_task_table(workbook)
    if candidate is None:
        raise GenericImportError(
            "未识别到可分析的任务表。普通进度表至少需要“任务名称”，并包含计划开始/完成日期或工期中的一种。"
        )
    ws, header_row, mapping, headers = candidate
    metadata = _extract_metadata(workbook)
    notes = [
        f"自动选择工作表“{ws.title}”第 {header_row} 行作为表头。",
        "字段映射：" + "；".join(f"{headers[column]} → {_field_label(field)}" for field, column in mapping.items()),
    ]
    refs: list[SourceRef] = []
    parsed_rows: list[dict[str, Any]] = []
    used_ids: set[str] = set()
    name_to_id: dict[str, str] = {}
    generated_ids = 0

    for row_index in range(header_row + 1, ws.max_row + 1):
        values = {field: ws.cell(row_index, column).value for field, column in mapping.items()}
        name = _text(values.get("name"))
        if not name:
            continue
        raw_wbs = values.get("wbs")
        wbs = _identifier(raw_wbs)
        if not wbs:
            generated_ids += 1
            wbs = f"AUTO-{generated_ids:03d}"
        base_wbs = wbs
        suffix = 2
        while wbs in used_ids:
            wbs = f"{base_wbs}-{suffix}"
            suffix += 1
        used_ids.add(wbs)
        name_to_id[_norm(name)] = wbs
        for field, column in mapping.items():
            value = ws.cell(row_index, column).value
            if value not in (None, ""):
                refs.append(SourceRef("自动导入", wbs, _field_label(field), ws.title, ws.cell(row_index, column).coordinate, value))
        parsed_rows.append({"row": row_index, "wbs": wbs, "generated_wbs": raw_wbs in (None, ""), "name": name, **values})

    if not parsed_rows:
        raise GenericImportError("已识别表头，但没有读取到包含任务名称的数据行。")

    raw_dates = []
    for row in parsed_rows:
        for field in ("plan_start", "plan_finish", "actual_start", "actual_finish"):
            value = _parse_date(row.get(field), workbook.epoch)
            if value:
                raw_dates.append(value)
    if not raw_dates and "duration" not in mapping:
        raise GenericImportError("任务表缺少可识别的计划日期和工期，无法生成横道图。")

    default_start = min(raw_dates) if raw_dates else date.today()
    calendar_spec = CalendarSpec("AUTO-CAL", {0, 1, 2, 3, 4})
    calendar = WorkCalendar(calendar_spec)
    tasks: list[Task] = []
    actuals: list[ActualProgress] = []
    baselines: list[BaselineRecord] = []
    relations: list[Relation] = []
    inferred_duration_count = 0
    inferred_progress_count = 0

    for index, row in enumerate(parsed_rows, start=1):
        start = _parse_date(row.get("plan_start"), workbook.epoch)
        finish = _parse_date(row.get("plan_finish"), workbook.epoch)
        duration = _parse_int(row.get("duration"))
        if duration is None and start and finish:
            duration = max(1, calendar.working_days_inclusive(start, finish))
        elif duration is not None:
            duration = max(1, duration)
        if start is None and finish and duration:
            start = calendar.add_workdays(finish, -(duration - 1))
            inferred_duration_count += 1
        if finish is None and start and duration:
            finish = calendar.add_workdays(start, duration - 1)
            inferred_duration_count += 1
        if start is None:
            start = default_start
            inferred_duration_count += 1
        if duration is None:
            duration = 1
            inferred_duration_count += 1
        if finish is None:
            finish = calendar.add_workdays(start, duration - 1)

        progress = _parse_percent(row.get("progress"))
        actual_start = _parse_date(row.get("actual_start"), workbook.epoch)
        actual_finish = _parse_date(row.get("actual_finish"), workbook.epoch)
        status = _status(row.get("status"), progress, actual_start, actual_finish)
        if progress is None:
            progress = 100.0 if status == "已完成" else 0.0
        quantity = _parse_float(row.get("quantity"))
        actual_quantity = _parse_float(row.get("actual_quantity"))
        if quantity is None or quantity <= 0:
            quantity = 100.0
            if actual_quantity is None:
                actual_quantity = progress
                inferred_progress_count += 1
            unit = _text(row.get("unit")) or "%"
        else:
            unit = _text(row.get("unit")) or "项"
            if actual_quantity is None:
                actual_quantity = quantity * progress / 100.0
                inferred_progress_count += 1
        actual_quantity = min(quantity, max(0.0, actual_quantity or 0.0))
        remaining = _parse_int(row.get("remaining"))
        if status == "进行中" and not remaining:
            remaining = max(1, math.ceil(duration * (1 - progress / 100.0)))
            inferred_progress_count += 1

        tasks.append(Task(
            wbs=row["wbs"],
            name=row["name"],
            discipline=_text(row.get("discipline")) or "待确认专业",
            section=_text(row.get("section")) or "待确认区段",
            station=_text(row.get("station")),
            workfront=_text(row.get("workfront")) or "待确认工作面",
            unit=unit,
            quantity=quantity,
            status=status,
            confirmed=not row["generated_wbs"],
            duration_override=duration,
            is_milestone=duration == 1 and ("里程碑" in row["name"] or "节点" in row["name"]),
        ))
        baselines.append(BaselineRecord(row["wbs"], None, None, start, finish))
        if status != "未开始" or progress > 0:
            actuals.append(ActualProgress(
                row["wbs"], actual_quantity, actual_start, actual_finish,
                0 if status == "已完成" else remaining,
                f"自动导入:{ws.title}!{row['row']}", "自动识别",
            ))

    for row in parsed_rows:
        predecessor = row.get("predecessor")
        if predecessor in (None, ""):
            continue
        for pred_id, relation_type, lag in _parse_predecessors(predecessor, name_to_id):
            if pred_id in used_ids and pred_id != row["wbs"]:
                relations.append(Relation(pred_id, row["wbs"], relation_type, lag, True))

    starts = [item.current_start for item in baselines if item.current_start]
    finishes = [item.current_finish for item in baselines if item.current_finish]
    project_start = min(starts) if starts else default_start
    planned_finish = max(finishes) if finishes else project_start
    contract_finish = _parse_date(metadata.get("contract_finish"), workbook.epoch) or planned_finish
    data_date = _parse_date(metadata.get("data_date"), workbook.epoch) or min(date.today(), max(project_start, planned_finish))
    project_id = _text(metadata.get("project_id")) or f"AUTO-{hashlib.sha256(content).hexdigest()[:8].upper()}"
    project_name = _text(metadata.get("project_name")) or Path(filename).stem
    project = Project(project_id, project_name, project_start, contract_finish, data_date, "AUTO-CAL", planned_finish)

    latest = max(baselines, key=lambda item: item.current_finish or project_start)
    milestones = [Milestone("AUTO-MS-01", "计划完工（自动识别）", latest.wbs, contract_finish, planned_finish)]
    if generated_ids:
        notes.append(f"{generated_ids} 项任务缺少 WBS，已生成 AUTO 编码，需人工确认。")
    if inferred_duration_count:
        notes.append(f"{inferred_duration_count} 处缺失日期/工期已按工作日历互推或采用 1 个工作日，需人工确认。")
    if inferred_progress_count:
        notes.append(f"{inferred_progress_count} 处工程量、完成率或剩余工期已按百分比/线性剩余工期换算，需人工确认。")
    if not relations:
        notes.append("未识别到有效前置关系；横道图和日期偏差可用，但关键线路仅为按完工日期识别的候选结果。")
    notes.append("未识别到企业资源定额；系统沿用原表工期，不生成量化资源赶工结论。")
    notes.append("默认采用周一至周五工作日历，未自动假定节假日和补班日期。")
    if "contract_finish" not in metadata:
        notes.append("未识别到合同竣工日，暂以原表最晚计划完成日作为合同日期。")
    if "data_date" not in metadata:
        notes.append(f"未识别到数据日期，暂采用 {data_date.isoformat()}。")

    return InputBundle(
        project=project,
        calendar=calendar_spec,
        tasks=tasks,
        relations=relations,
        resources=[],
        crash_params=[],
        constraints=[],
        actuals=actuals,
        delay_causes=[],
        period_plans=[],
        baselines=baselines,
        milestones=milestones,
        source_refs=refs,
        input_hash=hashlib.sha256(content).hexdigest(),
        input_filename=filename,
        import_mode="普通进度表自动映射",
        import_notes=notes,
    )


def _find_task_table(workbook):
    best = None
    for ws in workbook.worksheets:
        max_row = min(ws.max_row, 30)
        max_col = min(ws.max_column, 80)
        for row in range(1, max_row + 1):
            headers = {column: _text(ws.cell(row, column).value) for column in range(1, max_col + 1)}
            mapping = _map_headers(headers)
            if "name" not in mapping:
                continue
            has_timing = any(field in mapping for field in ("plan_start", "plan_finish", "duration"))
            if not has_timing:
                continue
            score = len(mapping) * 2 + sum(field in mapping for field in ("wbs", "plan_start", "plan_finish", "duration", "progress", "predecessor"))
            candidate = (score, ws, row, mapping, headers)
            if best is None or score > best[0]:
                best = candidate
    return None if best is None else best[1:]


def _map_headers(headers: dict[int, str]) -> dict[str, int]:
    result = {}
    normalized_aliases = {field: [_norm(alias) for alias in aliases] for field, aliases in ALIASES.items()}
    for column, header in headers.items():
        normalized = _norm(header)
        if not normalized:
            continue
        exact = [(field, aliases.index(normalized)) for field, aliases in normalized_aliases.items() if normalized in aliases]
        if exact:
            field = sorted(exact, key=lambda item: item[1])[0][0]
            result.setdefault(field, column)
            continue
        partial = [field for field, aliases in normalized_aliases.items() if any(len(alias) >= 3 and alias in normalized for alias in aliases)]
        if len(partial) == 1:
            result.setdefault(partial[0], column)
    return result


def _extract_metadata(workbook) -> dict[str, Any]:
    result = {}
    aliases = {field: {_norm(item) for item in values} for field, values in METADATA_ALIASES.items()}
    for ws in workbook.worksheets:
        for row in range(1, min(ws.max_row, 20) + 1):
            for column in range(1, min(ws.max_column, 12) + 1):
                key = _norm(ws.cell(row, column).value)
                for field, candidates in aliases.items():
                    if key in candidates and field not in result:
                        value = ws.cell(row, column + 1).value if column < ws.max_column else None
                        if value not in (None, ""):
                            result[field] = value
    return result


def _parse_predecessors(value: Any, name_to_id: dict[str, str]):
    result = []
    for token in re.split(r"[,，;；\n]+", _text(value)):
        token = token.strip()
        if not token:
            continue
        match = re.match(r"^(.*?)(?:\s*(FS|SS|FF|SF)\s*([+-]\s*\d+)?)?$", token, re.IGNORECASE)
        raw_id = (match.group(1) if match else token).strip()
        relation_type = (match.group(2) if match and match.group(2) else "FS").upper()
        lag_text = match.group(3).replace(" ", "") if match and match.group(3) else "0"
        pred_id = _identifier(raw_id)
        pred_id = name_to_id.get(_norm(raw_id), pred_id)
        result.append((pred_id, relation_type, int(lag_text)))
    return result


def _parse_date(value: Any, epoch) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = _text(value).replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        match = re.search(r"(20\d{2})[-](\d{1,2})[-](\d{1,2})", text)
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3))) if match else None


def _parse_percent(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, str):
        text = value.strip().replace("％", "%")
        try:
            number = float(text.rstrip("%"))
            return min(100.0, max(0.0, number if "%" in text else (number * 100 if 0 <= number <= 1 else number)))
        except ValueError:
            return None
    number = float(value)
    return min(100.0, max(0.0, number * 100 if 0 <= number <= 1 else number))


def _status(value: Any, progress: float | None, actual_start: date | None, actual_finish: date | None) -> str:
    text = _text(value)
    if actual_finish or progress == 100 or any(item in text for item in ("完成", "完工", "已结束")):
        return "已完成"
    if actual_start or (progress is not None and progress > 0) or any(item in text for item in ("进行", "施工中", "已开始")):
        return "进行中"
    return "未开始"


def _identifier(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _parse_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _parse_int(value: Any) -> int | None:
    number = _parse_float(value)
    return None if number is None else int(math.ceil(number))


def _text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _norm(value: Any) -> str:
    return re.sub(r"[\s_\-—（）()【】\[\]:：%％/]+", "", _text(value)).lower()


def _field_label(field: str) -> str:
    labels = {
        "wbs": "WBS编码", "name": "任务名称", "discipline": "专业", "section": "区段",
        "station": "桩号范围", "workfront": "工作面", "unit": "计量单位", "quantity": "总工程量",
        "actual_quantity": "实际完成工程量", "plan_start": "计划开始", "plan_finish": "计划完成",
        "duration": "工期", "predecessor": "前置任务", "progress": "完成率", "status": "任务状态",
        "actual_start": "实际开始", "actual_finish": "实际完成", "remaining": "剩余工期",
    }
    return labels.get(field, field)
